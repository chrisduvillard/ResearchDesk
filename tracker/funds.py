"""Official fund adapters; typed measures and conservative completeness checks."""

import csv
import hashlib
import io
import json
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from . import db
from .research import archive_bytes, activity
from .dbmf.markets import normalize_name, identify, EXPIRY, UnknownInstrument

VERSION = "funds-2"
# Explicit aliases, never fuzzy matches. Contracts retain their own instrument_key.
ALIASES = {
    "BRENT CRUDE FUTR": "brent",
    "COPPER FUTURE": "copper",
    "NY HARB ULSD FUT": "heating-oil",
    "KC HRW WHEAT FUT": "kc-wheat",
    "LIVE CATTLE FUTR": "live-cattle",
    "LOW SU GASOIL G": "gasoil",
    "WTI CRUDE FUTURE": "wti",
    "WHEAT FUTURE(CBT)": "wheat",
    "US 10YR ULTRA FUT": "us10ultra",
    "SUGAR #11 (WORLD)": "sugar",
    "SOYBEAN FUTURE": "soybeans",
    "SOYBEAN OIL FUTR": "soybean-oil",
    "NATURAL GAS FUTR": "natural-gas",
    "LONG GILT FUTURE": "uk10y",
    "JPN 10Y BOND(OSE)": "jp10y",
    "GASOLINE RBOB FUT": "gasoline",
    "EURO-BUND FUTURE": "de10y",
    "CORN FUTURE": "corn",
    "CHF CURRENCY FUT": "chf",
    "CAN 10YR BOND FUT": "ca10y",
    "C$ CURRENCY FUT": "cad",
    "BP CURRENCY FUT": "gbp",
    "AUDUSD CRNCY FUT": "aud",
}
ALIASES = {normalize_name(k): v for k, v in ALIASES.items()}
MEASURES = {
    "notional_pct_nav": "percent_nav",
    "collateral_pct_nav": "percent_nav",
    "equity_weight_pct": "percent_nav",
    "issuer_risk_weight_pct": "issuer_percent",
    "volatility_contribution_pct": "percent_volatility",
    "initial_margin": "USD",
    "issuer_portfolio_weight_pct": "issuer_percent",
}


def number(value):
    try:
        result = Decimal(
            str(value).replace(",", "").replace("$", "").replace("%", "").strip()
        )
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric field: {value!r}") from exc
    if not result.is_finite():
        raise ValueError("Nonfinite numeric field")
    return result


def market(name):
    if normalize_name(name) in ALIASES:
        return "market:" + ALIASES[normalize_name(name)]
    return "market:" + identify(name)


def expiry(name):
    match = EXPIRY.search(name)
    if not match:
        raise ValueError("Futures contract expiry missing: " + name)
    month = datetime.strptime(match[1].upper(), "%b").month
    year = int(match[2])
    year += 2000 if year < 100 else 0
    return f"{year:04}-{month:02}"


def holding(
    name,
    key,
    asset,
    measure,
    value,
    quantity=None,
    notional=None,
    expires=None,
    currency=None,
    evidence=None,
):
    return dict(
        name=name,
        instrument_key=key,
        asset_id=asset,
        measure=measure,
        value=str(number(value)),
        unit=MEASURES[measure],
        quantity=str(number(quantity)) if quantity is not None else None,
        notional=str(number(notional)) if notional is not None else None,
        expiry=expires,
        currency=currency,
        collateral=measure == "collateral_pct_nav",
        evidence=evidence or name,
    )


def parse_kmlm(raw, *, net_assets, expected_date, expected_rows=None):
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    if len(rows) < 3 or rows[0][:2] != ["KMLM Holdings", "As of " + expected_date]:
        raise ValueError("KMLM identity or source date mismatch")
    nav = number(net_assets)
    if nav <= 0:
        raise ValueError("Positive NAV required")
    fields = [
        "Rank",
        "Company Name",
        "% of Net Assets",
        "Ticker",
        "Identifier",
        "Type",
        "Shares Held",
        "Market Value($)",
        "Notional Value($)",
    ]
    if rows[1] != fields:
        raise ValueError("KMLM columns changed")
    records = [dict(zip(fields, r)) for r in rows[2:] if any(r)]
    if expected_rows is not None and len(records) != expected_rows:
        raise ValueError("KMLM incomplete row count")
    if len(records) < 2 or [int(r["Rank"]) for r in records] != list(
        range(1, len(records) + 1)
    ):
        raise ValueError("KMLM missing or duplicate ranks")
    holdings = []
    collateral = sum(
        (
            number(r["Market Value($)"])
            for r in records
            if r["Type"] in ("Cash", "Currency", "Treasury Bill")
        ),
        Decimal(0),
    )
    if collateral <= 0:
        raise ValueError("KMLM collateral total must be positive")
    collateral_rows = [
        r for r in records if r["Type"] in ("Cash", "Currency", "Treasury Bill")
    ]
    if not any(
        all(
            abs(
                number(r["% of Net Assets"])
                - number(r["Market Value($)"]) / basis * 100
            )
            <= Decimal(".02")
            for r in collateral_rows
        )
        for basis in (nav, collateral)
    ):
        raise ValueError("KMLM reported collateral weight mismatch")
    for r in records:
        name = r["Company Name"]
        kind = r["Type"]
        q = number(r["Shares Held"])
        n = number(r["Notional Value($)"])
        if kind == "Future":
            if q == 0 or n == 0 or (q > 0) != (n > 0):
                raise ValueError("KMLM quantity/notional sign mismatch")
            holdings.append(
                holding(
                    name,
                    r["Ticker"] + ":" + name,
                    market(name),
                    "notional_pct_nav",
                    n / nav * 100,
                    q,
                    n,
                    expiry(name),
                    "USD",
                    json.dumps(r),
                )
            )
        elif kind in ("Cash", "Currency", "Treasury Bill"):
            value = number(r["% of Net Assets"])
            mv = number(r["Market Value($)"])
            # The issuer's displayed weights can use a denominator different
            # from its dated NAV. Validate those weights against the complete
            # collateral basket; retain them separately from calculated NAV ratios.
            asset = (
                "market:tbills"
                if kind == "Treasury Bill"
                else "market:cash"
                if kind == "Cash"
                else None
            )
            holdings.append(
                holding(
                    name,
                    r["Identifier"] + ":" + name,
                    asset,
                    "collateral_pct_nav",
                    mv / nav * 100,
                    q,
                    mv,
                    None,
                    "USD",
                    json.dumps(r),
                )
            )
            holdings.append(
                holding(
                    name,
                    r["Identifier"] + ":" + name,
                    asset,
                    "issuer_portfolio_weight_pct",
                    value,
                    evidence=json.dumps(r),
                )
            )
        else:
            raise ValueError("KMLM unfamiliar security type: " + kind)
    if abs(collateral / nav - 1) > Decimal(".02"):
        raise ValueError("KMLM incomplete collateral total")
    return dict(
        source_date=expected_date,
        complete=True,
        holdings=holdings,
        metadata={
            "net_assets": str(nav),
            "rows": len(records),
            "collateral_market_value": str(collateral),
            "reported_weight_basis": "Issuer weights retained separately; calculated exposures use the published dated NAV",
        },
    )


def parse_cta(raw, *, expected_date):
    from zipfile import ZipFile

    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 50_000_000:
            raise ValueError("Workbook exceeds expanded size limit")
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        rows = book["Simplify Portfolio Tracker"].iter_rows(values_only=True)
        top = next(rows)
        stamp = datetime.strptime(str(top[6]), "%m/%d/%Y").date().isoformat()
        if stamp != expected_date:
            raise ValueError("CTA workbook source date mismatch")
        fields = next(rows)
        records = [dict(zip(fields, r)) for r in rows if r[0] == "CTA"]
        if not records:
            raise ValueError("CTA missing from workbook")
        holdings = []
        collateral = Decimal(0)
        navs = set()
        for r in records:
            name = r["SECURITY DESCRIPTION"]
            kind = r["Security Master Asset Group"]
            nav = number(r["Total MV"])
            navs.add(nav)
            value = number(r["Weight"]) * 100
            exposure = number(r["Market Value/Exposure"])
            q = number(r["Quantity"])
            official_total = number(r["Official NAV"]) * number(r["Shares Out"])
            projected_total = number(r["BNY Projected NAV"]) * number(r["Shares Out"])
            if (
                nav <= 0
                or projected_total <= 0
                or abs(nav - official_total) > Decimal(".01")
                or abs(value - exposure / projected_total * 100) > Decimal(".00001")
            ):
                raise ValueError("CTA inconsistent units/NAV")
            reported_weight = value
            value = exposure / nav * 100
            if kind == "Future":
                if q == 0 or exposure == 0 or (q > 0) != (exposure > 0):
                    raise ValueError("CTA futures sign mismatch")
                row = holding(
                    name,
                    str(r["ICE Ticker/Identifier"]),
                    market(name),
                    "notional_pct_nav",
                    value,
                    q,
                    exposure,
                    expiry(name),
                    "USD",
                )
            elif (
                kind in ("Cash", "Treasury Bill")
                or kind == "Fund"
                and r["TICKER"] == "SBIL"
            ):
                collateral += exposure
                asset = (
                    "market:cash"
                    if kind == "Cash"
                    else "market:tbills"
                    if kind == "Treasury Bill"
                    else None
                )
                row = holding(
                    name,
                    str(r["ISIN"] or r["ICE Ticker/Identifier"]),
                    asset,
                    "collateral_pct_nav",
                    value,
                    q,
                    exposure,
                    None,
                    "USD",
                )
            else:
                raise ValueError("CTA unfamiliar security: " + name)
            row["evidence"] = json.dumps(r, default=str)
            holdings.append(row)
            holdings.append(
                holding(
                    name,
                    row["instrument_key"],
                    row["asset_id"],
                    "issuer_portfolio_weight_pct",
                    reported_weight,
                    evidence="Issuer supplied Weight column; denominator differs from Total MV",
                )
            )
        if len(navs) != 1 or abs(collateral / nav - 1) > Decimal(".02"):
            raise ValueError("CTA incomplete collateral or inconsistent NAV")
        risk = book["CTA Est. Risk Profile"].iter_rows(values_only=True)
        if next(risk)[0].date().isoformat() != expected_date:
            raise ValueError("CTA risk profile date mismatch")
        if tuple(next(risk)[:4]) != (
            "Category",
            "Weight*",
            "Est. Initial Margin",
            "Contrib to Vol",
        ):
            raise ValueError("CTA risk columns changed")
        risk_total = Decimal(0)
        for name, weight, margin, vol in risk:
            if name == "Total":
                break
            asset = market(name)
            risk_total += number(vol)
            for measure, value in [
                ("issuer_risk_weight_pct", number(weight) * 100),
                ("initial_margin", number(margin)),
                ("volatility_contribution_pct", number(vol) * 100),
            ]:
                holdings.append(
                    holding(
                        name,
                        "risk:" + name,
                        asset,
                        measure,
                        value,
                        evidence="Issuer risk profile; bond weights use 10-year equivalents",
                    )
                )
        if abs(risk_total - 1) > Decimal(".001"):
            raise ValueError("CTA incomplete risk profile")
        futures = {
            r["asset_id"] for r in holdings if r["measure"] == "notional_pct_nav"
        }
        risk_assets = {
            r["asset_id"] for r in holdings if r["measure"] == "issuer_risk_weight_pct"
        }
        if futures != risk_assets:
            raise ValueError("CTA holdings and risk profile markets differ")
        return dict(
            source_date=stamp,
            complete=True,
            holdings=holdings,
            metadata={
                "net_assets": str(nav),
                "reported_weight_basis": "BNY Projected NAV times Shares Out; computed notional/NAV uses Official NAV times Shares Out",
                "risk_basis": "Issuer weights: 10-year equivalents for interest-rate and bond futures",
            },
        )
    finally:
        book.close()


def qualify(conn, source_id):
    rows = conn.execute(
        "SELECT * FROM source_runs WHERE source_id=? AND scheduled=1 ORDER BY id DESC",
        (source_id,),
    ).fetchall()
    slots = []
    seen = set()
    for r in rows:
        if r["slot"] in seen:
            continue
        seen.add(r["slot"])
        slots.append(r)
        if len(slots) == 5:
            break
    return (
        len(slots) == 5
        and all(r["status"] == "success" for r in slots)
        and len({r["report_date"] for r in slots}) >= 2
    )


def ensure_market(conn, asset, name):
    if not asset:
        return
    conn.execute(
        "INSERT OR IGNORE INTO assets(id,symbol,name,kind,verified,provenance) VALUES(?,?,?,?,?,?)",
        (
            asset,
            asset.split(":", 1)[-1],
            name,
            "futures_market" if asset.startswith("market:") else "security",
            int(asset.startswith("market:") or asset.startswith("cusip:")),
            "Official fund holding; adapter " + VERSION,
        ),
    )


def accept_report(conn, fund, document, raw, acquired=None):
    acquired = acquired or db.utcnow()
    if not document["complete"] or not document["holdings"]:
        raise ValueError("Incomplete fund report")
    day = date.fromisoformat(document["source_date"])
    if day > acquired.date():
        raise ValueError("Fund source date is in the future")
    keys = set()
    for row in document["holdings"]:
        key = (row["instrument_key"], row["measure"])
        if key in keys:
            raise ValueError("Duplicate holding")
        keys.add(key)
        if MEASURES.get(row["measure"]) != row["unit"]:
            raise ValueError("Incompatible measure units")
        number(row["value"])
    # Fingerprint the parsed content, not whitespace or provider page decoration.
    normalized = dict(
        document,
        holdings=sorted(
            document["holdings"], key=lambda r: (r["instrument_key"], r["measure"])
        ),
    )
    fingerprint = hashlib.sha256(
        json.dumps(normalized, sort_keys=True).encode()
    ).hexdigest()
    identifier = f"{fund}:{fingerprint}"
    path, digest = archive_bytes(raw, "funds")
    source = "fund:" + fund
    with conn:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute(
            "SELECT * FROM fund_reports WHERE id=?", (identifier,)
        ).fetchone()
        if not existing:
            previous = conn.execute(
                "SELECT * FROM fund_reports WHERE fund_id=? AND status='accepted' ORDER BY source_date DESC,revision DESC LIMIT 1",
                (fund,),
            ).fetchone()
            if previous and document["source_date"] < previous["source_date"]:
                raise ValueError("Older fund report; current state preserved")
            revision = conn.execute(
                "SELECT coalesce(max(revision),-1)+1 FROM fund_reports WHERE fund_id=? AND source_date=?",
                (fund, document["source_date"]),
            ).fetchone()[0]
            conn.execute(
                """INSERT INTO fund_reports(id,fund_id,source_id,source_date,acquired_at,processed_at,raw_path,content_hash,parser_version,status,revision,complete,metadata_json)
              VALUES(?,?,?,?,?,?,?,?,?,'accepted',?,1,?)""",
                (
                    identifier,
                    fund,
                    source,
                    document["source_date"],
                    db.iso(acquired),
                    db.iso(),
                    path,
                    digest,
                    VERSION,
                    revision,
                    json.dumps(document["metadata"]),
                ),
            )
            for r in normalized["holdings"]:
                ensure_market(conn, r["asset_id"], r["name"])
                if r.get("display_symbol"):
                    conn.execute(
                        "UPDATE assets SET symbol=? WHERE id=?",
                        (r["display_symbol"], r["asset_id"]),
                    )
                fields = [
                    "asset_id",
                    "instrument_key",
                    "name",
                    "measure",
                    "value",
                    "unit",
                    "quantity",
                    "notional",
                    "expiry",
                    "currency",
                    "collateral",
                    "evidence",
                ]
                conn.execute(
                    f"INSERT INTO fund_holdings(report_id,{','.join(fields)}) VALUES({','.join('?' for _ in range(len(fields) + 1))})",
                    (identifier, *(r.get(k) for k in fields)),
                )
            activity(
                conn,
                "fund:" + identifier,
                source,
                None,
                "correction" if revision else "fund_report",
                document["source_date"],
                db.iso(acquired),
                f"{fund}: " + ("corrected report" if revision else "holdings report"),
                "/api/v2/fund-reports/" + identifier + "/source",
                "fund_report",
                identifier,
            )
            publish_changes(conn, identifier)
        conn.execute(
            "INSERT INTO fund_observations(report_id,acquired_at,raw_path,content_hash) VALUES(?,?,?,?)",
            (identifier, db.iso(acquired), path, digest),
        )
    return dict(
        conn.execute("SELECT * FROM fund_reports WHERE id=?", (identifier,)).fetchone()
    )


def sync_dbmf(conn):
    with conn:
        for m in conn.execute("SELECT * FROM dbmf_markets").fetchall():
            ensure_market(conn, "market:" + m["id"], m["name"])
        for r in conn.execute(
            "SELECT * FROM dbmf_reports WHERE status='accepted' AND id NOT IN (SELECT legacy_id FROM fund_reports WHERE legacy_id IS NOT NULL)"
        ).fetchall():
            identifier = "DBMF:" + str(r["id"])
            acquired = conn.execute(
                "SELECT min(collected_at) FROM dbmf_observations WHERE report_id=? AND timestamp_basis='acquisition'",
                (r["id"],),
            ).fetchone()[0]
            metadata = json.loads(r["metadata_json"])
            metadata.update(
                source_kind=r["source_kind"],
                legacy_revision=r["revision"],
                replayed_from=r["replayed_from"],
                timestamp_basis="acquisition" if acquired else "unknown",
            )
            conn.execute(
                """INSERT INTO fund_reports(id,fund_id,source_id,source_date,acquired_at,processed_at,raw_path,content_hash,parser_version,status,revision,complete,metadata_json,legacy_id)
              VALUES(?,'DBMF','fund:DBMF',?,?,?,?,?,?,'accepted',?,1,?,?)""",
                (
                    identifier,
                    r["source_date"],
                    acquired,
                    r["processed_at"] or db.iso(),
                    r["raw_path"],
                    r["raw_hash"],
                    r["parser_version"],
                    max(0, r["revision"] - 1),
                    json.dumps(metadata),
                    r["id"],
                ),
            )
            occurrences = defaultdict(int)
            for h in conn.execute(
                "SELECT h.*,m.category FROM dbmf_holdings h JOIN dbmf_markets m ON m.id=h.market_id WHERE report_id=?",
                (r["id"],),
            ).fetchall():
                measure = (
                    "collateral_pct_nav"
                    if h["category"] == "Collateral"
                    else "notional_pct_nav"
                )
                contract = "|".join(
                    str(h[k] or "")
                    for k in (
                        "market_id",
                        "identifier",
                        "ticker",
                        "expiry",
                        "original_name",
                    )
                )
                occurrences[contract] += 1
                instrument_key = contract + "::lot:" + str(occurrences[contract])
                conn.execute(
                    """INSERT INTO fund_holdings(report_id,asset_id,instrument_key,name,measure,value,unit,quantity,notional,expiry,currency,collateral,evidence)
                  VALUES(?,?,?,?,?,?,'percent_nav',?,?,?,'USD',?,?)""",
                    (
                        identifier,
                        "market:" + h["market_id"],
                        instrument_key,
                        h["original_name"],
                        measure,
                        h["exposure_pct"],
                        h["quantity"],
                        h["notional"],
                        h["expiry"],
                        int(measure == "collateral_pct_nav"),
                        h["evidence"],
                    ),
                )
            activity(
                conn,
                "fund:" + identifier,
                "fund:DBMF",
                None,
                "historical_report"
                if not acquired or r["source_kind"] != "live"
                else "correction"
                if r["revision"] > 1 or r["replayed_from"]
                else "fund_report",
                r["source_date"],
                acquired or r["processed_at"] or db.iso(),
                f"DBMF: report {r['source_date']}",
                "/api/v2/fund-reports/" + identifier + "/source",
                "fund_report",
                identifier,
            )
            publish_changes(conn, identifier)
        # Preserve the established live-over-import and repeated-observation selection.
        from .dbmf.store import reports as legacy_reports

        canonical = [r["id"] for r in legacy_reports(conn)]
        conn.execute(
            "UPDATE fund_reports SET status=CASE WHEN legacy_id IN (SELECT value FROM json_each(?)) THEN 'accepted' ELSE 'archived' END WHERE fund_id='DBMF'",
            (json.dumps(canonical),),
        )
        conn.execute("""INSERT INTO fund_observations(report_id,acquired_at,raw_path,content_hash,legacy_id)
          SELECT 'DBMF:'||o.report_id,o.collected_at,o.raw_path,o.raw_hash,o.id FROM dbmf_observations o
          JOIN fund_reports f ON f.legacy_id=o.report_id WHERE o.timestamp_basis='acquisition' AND o.id NOT IN
          (SELECT legacy_id FROM fund_observations WHERE legacy_id IS NOT NULL)""")


def comparison(conn, funds, measure, aligned=False):
    if measure not in MEASURES:
        raise ValueError("Unknown measure")
    byfund = {
        f: conn.execute(
            "SELECT * FROM fund_reports WHERE fund_id=? AND complete=1 AND status='accepted' ORDER BY source_date DESC,revision DESC",
            (f,),
        ).fetchall()
        for f in funds
    }
    dates = None
    if aligned:
        for reports in byfund.values():
            found = {r["source_date"] for r in reports}
            dates = found if dates is None else dates & found
    selected = {
        f: next(
            (
                r
                for r in reports
                if not aligned or dates and r["source_date"] == max(dates)
            ),
            None,
        )
        for f, reports in byfund.items()
    }
    columns = set()
    rows = []
    for fund, report in selected.items():
        cells = {}
        if report:
            prior = next(
                (r for r in byfund[fund] if r["source_date"] < report["source_date"]),
                None,
            )

            def aggregate(r):
                result = defaultdict(Decimal)
                if r:
                    for h in conn.execute(
                        "SELECT * FROM fund_holdings WHERE report_id=? AND measure=?",
                        (r["id"], measure),
                    ):
                        result[h["asset_id"] or "unmapped:" + h["instrument_key"]] += (
                            number(h["value"])
                        )
                return result

            current, old = aggregate(report), aggregate(prior)
            for key in current.keys() | old.keys():
                value = current.get(key, Decimal(0))
                cells[key] = {
                    "value": float(value),
                    "change_pp": float(value - old.get(key, 0)) if prior else None,
                }
            columns.update(cells)
        rows.append(
            dict(
                fund_id=fund,
                report_id=report["id"] if report else None,
                source_date=report["source_date"] if report else None,
                stale=not report
                or (db.utcnow().date() - date.fromisoformat(report["source_date"])).days
                > 4,
                cells=cells,
            )
        )
    for row in rows:
        for key in columns:
            row["cells"].setdefault(key, {"value": None, "change_pp": None})
    return dict(
        measure=measure,
        unit=MEASURES[measure],
        mode="aligned_date" if aligned else "latest_available",
        columns=sorted(columns),
        rows=rows,
    )


def parse_ark(raw, fund):
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    required = [
        "date",
        "fund",
        "company",
        "ticker",
        "cusip",
        "shares",
        "market value ($)",
        "weight (%)",
    ]
    if reader.fieldnames != required:
        raise ValueError("ARK holdings columns changed")
    holdings = []
    dates = set()
    total = Decimal(0)
    seen = set()
    footer = False
    for r in reader:
        if not r["date"] or not r["date"].strip():
            continue
        if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", r["date"]):
            # Only the published legal footer is allowed after the full table.
            if not r["date"].startswith(
                (
                    "Investors should",
                    "The principal",
                    "This and other",
                    "Foreside",
                    "ARK Investment",
                    "Holdings are",
                    "Please note",
                )
            ):
                raise ValueError("Unrecognized ARK trailing row")
            footer = True
            continue
        if footer or r["fund"] != fund or None in r:
            raise ValueError("ARK fund identity or row structure mismatch")
        dates.add(datetime.strptime(r["date"], "%m/%d/%Y").date().isoformat())
        if (
            not re.fullmatch(r"(?:[A-Z0-9]{9}|[A-Z0-9]{7})", r["cusip"] or "")
            or r["cusip"] in seen
        ):
            raise ValueError("ARK missing or duplicate security identifier")
        seen.add(r["cusip"])
        weight = number(r["weight (%)"])
        q = number(r["shares"])
        value = number(r["market value ($)"])
        if weight < 0 or q < 0 or value < 0:
            raise ValueError("Unexpected signed ARK equity holding")
        total += weight
        h = holding(
            r["company"],
            r["cusip"],
            ("cusip:" if len(r["cusip"]) == 9 else "ark-id:") + r["cusip"],
            "equity_weight_pct",
            weight,
            q,
            value,
            None,
            "USD",
            json.dumps(r),
        )
        h["display_symbol"] = r["ticker"]
        holdings.append(h)
    if len(dates) != 1 or len(holdings) < 2 or abs(total - 100) > Decimal("1"):
        raise ValueError("ARK incomplete holdings or mixed dates")
    return dict(
        source_date=dates.pop(),
        complete=True,
        holdings=holdings,
        metadata={
            "reported_weight_sum": str(total),
            "sector_coverage": "unavailable",
            "trade_notifications": "not inferred from holdings",
        },
    )


def report_changes(conn, report_id):
    report = conn.execute(
        "SELECT * FROM fund_reports WHERE id=?", (report_id,)
    ).fetchone()
    if not report:
        return []
    previous = conn.execute(
        "SELECT * FROM fund_reports WHERE fund_id=? AND status='accepted' AND (source_date<? OR source_date=? AND revision<?) ORDER BY source_date DESC,revision DESC LIMIT 1",
        (
            report["fund_id"],
            report["source_date"],
            report["source_date"],
            report["revision"],
        ),
    ).fetchone()
    if not previous:
        return []

    def grouped(identifier):
        values = {}
        for row in conn.execute(
            "SELECT * FROM fund_holdings WHERE report_id=?", (identifier,)
        ):
            key = (
                row["asset_id"] or "unmapped:" + row["instrument_key"],
                row["measure"],
            )
            if key not in values:
                values[key] = {
                    "value": Decimal(0),
                    "contracts": set(),
                    "name": row["name"],
                }
            values[key]["value"] += number(row["value"])
            values[key]["contracts"].add(row["instrument_key"].split("::lot:")[0])
        return values

    old, new = grouped(previous["id"]), grouped(report_id)
    result = []
    for key in sorted(old.keys() | new.keys()):
        before = old.get(key, {"value": Decimal(0), "contracts": set()})
        after = new.get(key, {"value": Decimal(0), "contracts": set()})
        rolled = key[1] == "notional_pct_nav" and bool(
            before["contracts"]
            and after["contracts"]
            and before["contracts"] != after["contracts"]
        )
        if before["value"] == after["value"] and not rolled:
            continue
        result.append(
            dict(
                asset_id=key[0] if not key[0].startswith("unmapped:") else None,
                measure=key[1],
                old=float(before["value"]),
                new=float(after["value"]),
                change_pp=float(after["value"] - before["value"]),
                contract_roll=rolled,
                kind="correction" if report["revision"] else "exposure_change",
                name=(new.get(key) or old[key])["name"],
                previous_report_id=previous["id"],
                report_id=report_id,
                source_date=report["source_date"],
            )
        )
    return result


def publish_changes(conn, report_id):
    report = conn.execute(
        "SELECT * FROM fund_reports WHERE id=?", (report_id,)
    ).fetchone()
    for item in report_changes(conn, report_id):
        metadata = json.loads(report["metadata_json"])
        if not report["acquired_at"] or metadata.get("source_kind", "live") != "live":
            item["kind"] = "historical_report"
        elif metadata.get("replayed_from"):
            item["kind"] = "correction"
        elif (
            datetime.fromisoformat(report["acquired_at"]).date()
            - date.fromisoformat(report["source_date"])
        ).days > 4:
            item["kind"] = "historical_report"
        key = hashlib.sha256(json.dumps(item, sort_keys=True).encode()).hexdigest()
        activity(
            conn,
            "fund-change:" + key,
            report["source_id"],
            item["asset_id"],
            item["kind"],
            report["source_date"],
            report["acquired_at"] or report["processed_at"],
            f"{report['fund_id']}: {item['name']} {item['change_pp']:+.2f} pp"
            + (" (contract roll)" if item["contract_roll"] else ""),
            "/api/v2/fund-reports/" + report_id + "/source",
            "fund_report",
            report_id,
            item,
        )
